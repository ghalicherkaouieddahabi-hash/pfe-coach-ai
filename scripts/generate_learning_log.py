"""Génère le Learning Log Word pour le rapport / le prof."""

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

ROWS = [
    (
        "1",
        "Interface web : Writer et Reviewer vides (« aucun contenu ») alors que Analyseur et Planificateur sont remplis.",
        "Décalage API (ancien format 3 agents : analyzer, planner, coach) vs UI qui lit pipeline.writer / pipeline.reviewer. Serveur non redémarré.",
        "Mise à jour extractPipeline() dans static/index.html ; payload API avec writer/reviewer ; agent writer_reviewer ; redémarrage Uvicorn + Ctrl+F5.",
        "static/index.html, services/pfe_pipeline.py, routes/analyze_route.py",
    ),
    (
        "2",
        "HTTP 429 — Quota exceeded, limit: 0 pour gemini-2.0-flash.",
        "Quota gratuit indisponible pour ce modèle ; fallbacks encore sur gemini-2.0-flash.",
        "GEMINI_MODEL=gemini-2.5-flash + fallbacks dans .env ; défauts dans gemini_service.py ; attente entre analyses.",
        "https://ai.google.dev/gemini-api/docs/rate-limits — https://ai.dev/rate-limit",
    ),
    (
        "3",
        "Échec ANALYSEUR / PLANIFICATEUR : « Tous les modèles Gemini configurés ont échoué ».",
        "Clé absente, quota atteint, ou .env non rechargé par Uvicorn.",
        "Vérifier GEMINI_API_KEY ; load_dotenv racine dans main.py ; GET / (gemini_model_chain).",
        "https://aistudio.google.com/apikey — README.md",
    ),
    (
        "4",
        "PowerShell : « && » invalide.",
        "Syntaxe bash utilisée dans PowerShell.",
        "start.ps1 ou Set-Location ; .\\venv\\Scripts\\uvicorn.exe ...",
        "Microsoft PowerShell — start.ps1 / start.bat",
    ),
    (
        "5",
        "Changement .env sans effet sur le modèle Gemini.",
        "Processus Uvicorn ancien ; reload ne recharge pas toujours l'environnement.",
        "Arrêt complet du serveur puis relance ; contrôle GET /.",
        "main.py — endpoint GET /",
    ),
    (
        "6",
        "Réponse API incomplète dans Swagger (pas de pipeline writer/reviewer).",
        "Ancienne version du pipeline encore en exécution.",
        "Redémarrer serveur ; vérifier pipeline.* dans la réponse JSON.",
        "http://127.0.0.1:8000/docs",
    ),
    (
        "7",
        "Analyse très lente ou timeout.",
        "~6 appels Gemini + pauses entre agents.",
        "Prévoir 2–5 min ; Writer+Reviewer fusionnés ; pas d'analyses parallèles.",
        "services/pfe_pipeline.py — README.md",
    ),
    (
        "8",
        "Export PDF/DOCX manquant ou 404.",
        "Échec export ou mauvais job_id dans l'URL.",
        "Vérifier pipeline.exports ; GET /exports/{job_id}/{file} ; pip install python-docx fpdf2.",
        "routes/export_route.py — export_service.py",
    ),
    (
        "9",
        "Confusion agents Python vs tools / MCP.",
        "Vocabulaire proche, technologies différentes.",
        "Documentation ; ajout serveur MCP (mcp_server/, .cursor/mcp.json).",
        "https://modelcontextprotocol.io — mcp_server/server.py",
    ),
    (
        "10",
        "Clé API Gemini exposée ou partagée.",
        ".env copié ou commité par erreur.",
        ".env dans .gitignore ; .env.example sans clé ; régénérer la clé si fuite.",
        ".gitignore — Google AI Studio",
    ),
]


def main() -> None:
    doc = Document()

    title = doc.add_heading("Journal d'apprentissage (Learning Log)", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph("Projet : Coach-PFE — Assistant IA de rédaction académique (PFE)")
    doc.add_paragraph("Stack : Python 3.12, FastAPI, Google Gemini, RAG, export PDF/DOCX, MCP")
    doc.add_paragraph(f"Date de rédaction : {date.today().strftime('%d/%m/%Y')}")
    doc.add_paragraph("Étudiant(s) : [Nom Prénom — à compléter]")
    doc.add_paragraph("Encadrant : [Nom du professeur — à compléter]")

    doc.add_heading("1. Objectif", level=1)
    doc.add_paragraph(
        "Documenter les erreurs rencontrées pendant le développement, leurs causes, "
        "les solutions mises en œuvre et les sources consultées (learning log exigé par l'encadrement)."
    )

    doc.add_heading("2. Tableau synthétique", level=1)

    headers = ("N°", "Erreur / symptôme", "Cause", "Solution", "Source")
    table = doc.add_table(rows=1 + len(ROWS), cols=5)
    table.style = "Table Grid"

    for j, h in enumerate(headers):
        table.rows[0].cells[j].text = h
        for run in table.rows[0].cells[j].paragraphs[0].runs:
            run.bold = True

    for i, row_data in enumerate(ROWS, start=1):
        for j, text in enumerate(row_data):
            table.rows[i].cells[j].text = text

    doc.add_heading("3. Leçons apprises", level=1)
    for item in [
        "Aligner le contrat API et le frontend avant les tests utilisateur.",
        "Anticiper les quotas API (modèle, fallbacks, nombre d'appels).",
        "Redémarrer le serveur après modification du .env ou de l'architecture.",
        "Tenir un journal d'erreurs en continu pour le rapport et la soutenance.",
    ]:
        doc.add_paragraph(item, style="List Bullet")

    doc.add_heading("4. Améliorations futures", level=1)
    doc.add_paragraph(
        "File d'attente asynchrone, cache RAG, monitoring quota, authentification, déploiement Docker."
    )

    doc.add_heading("5. Références", level=1)
    for ref in [
        "Google AI — Gemini rate limits : https://ai.google.dev/gemini-api/docs/rate-limits",
        "FastAPI : https://fastapi.tiangolo.com/",
        "Model Context Protocol : https://modelcontextprotocol.io",
        "Dépôt Coach-PFE — README.md",
    ]:
        doc.add_paragraph(ref, style="List Bullet")

    doc.add_paragraph()
    doc.add_paragraph("Signature étudiant : _________________________    Date : _____________")

    out = Path(__file__).resolve().parent.parent / "docs" / "Learning_Log_Coach-PFE.docx"
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out))
    print(f"Fichier créé : {out}")


if __name__ == "__main__":
    main()
