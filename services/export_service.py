"""
Export PDF / DOCX du document final généré par le pipeline.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path

logger = logging.getLogger(__name__)

_EXPORT_ROOT = Path(__file__).resolve().parent.parent / "exports"


def _safe_filename(topic: str) -> str:
    slug = re.sub(r"[^\w\s-]", "", topic, flags=re.UNICODE)[:40].strip().replace(" ", "_")
    return slug or "pfe"


def export_documents(job_id: str, topic: str, final_text: str) -> dict[str, str]:
    """
    Crée les fichiers DOCX et PDF dans exports/{job_id}/.

    :return: chemins relatifs pour téléchargement API, ex. ``exports/abc/doc.docx``.
    """
    out_dir = _EXPORT_ROOT / job_id
    out_dir.mkdir(parents=True, exist_ok=True)
    base = _safe_filename(topic)
    docx_path = out_dir / f"{base}.docx"
    pdf_path = out_dir / f"{base}.pdf"

    _write_docx(docx_path, topic, final_text)
    _write_pdf(pdf_path, topic, final_text)

    return {
        "docx": f"exports/{job_id}/{docx_path.name}",
        "pdf": f"exports/{job_id}/{pdf_path.name}",
    }


def _write_docx(path: Path, title: str, body: str) -> None:
    from docx import Document  # type: ignore[import-untyped]

    doc = Document()
    doc.add_heading(title[:200], level=0)
    for block in body.split("\n\n"):
        block = block.strip()
        if not block:
            continue
        if block.startswith("## "):
            doc.add_heading(block[3:].strip(), level=1)
        elif block.startswith("### "):
            doc.add_heading(block[4:].strip(), level=2)
        elif block.startswith("- "):
            for line in block.split("\n"):
                if line.strip().startswith("- "):
                    doc.add_paragraph(line.strip()[2:], style="List Bullet")
        else:
            doc.add_paragraph(block)
    doc.save(str(path))
    logger.info("Export DOCX : %s", path)


def _write_pdf(path: Path, title: str, body: str) -> None:
    from fpdf import FPDF  # type: ignore[import-untyped]

    pdf = FPDF()
    pdf.set_auto_page_break(auto=True, margin=15)
    pdf.add_page()
    pdf.set_font("Helvetica", "B", 14)
    pdf.multi_cell(0, 8, title[:200])
    pdf.ln(4)
    pdf.set_font("Helvetica", size=10)
    for line in body.split("\n"):
        line = line.strip()
        if not line:
            pdf.ln(3)
            continue
        if line.startswith("## "):
            pdf.set_font("Helvetica", "B", 12)
            pdf.multi_cell(0, 6, line[3:])
            pdf.set_font("Helvetica", size=10)
        else:
            safe = line.encode("latin-1", errors="replace").decode("latin-1")
            pdf.multi_cell(0, 5, safe)
    pdf.output(str(path))
    logger.info("Export PDF : %s", path)


def resolve_export_path(job_id: str, filename: str) -> Path | None:
    """Résout ``exports/{job_id}/{filename}`` de façon sécurisée."""
    safe_name = Path(filename).name
    if safe_name != filename or ".." in filename:
        return None
    full = (_EXPORT_ROOT / job_id / safe_name).resolve()
    root = _EXPORT_ROOT.resolve()
    if not str(full).startswith(str(root)):
        return None
    if full.is_file():
        return full
    return None
